import copy
from pptx import Presentation
from pptx.oxml.ns import qn
orig = Presentation('orig.pptx'); new = Presentation('new.pptx')
TOTAL = 18 + len(new.slides)
# 既存スライドの総ページ表記を更新
for s in orig.slides:
    for sh in s.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.text.strip().endswith('/ 18'):
                        r.text = r.text.replace('/ 18', f'/ {TOTAL}')
layout = orig.slide_layouts[0]
for s in new.slides:
    ns = orig.slides.add_slide(layout)
    for ph in list(ns.placeholders):
        ph._element.getparent().remove(ph._element)
    src_csld = s._element.find(qn('p:cSld')); dst_csld = ns._element.find(qn('p:cSld'))
    bg = src_csld.find(qn('p:bg'))
    if bg is not None:
        dst_csld.insert(0, copy.deepcopy(bg))
    tree = ns.shapes._spTree
    for el in s.shapes._spTree:
        if el.tag in (qn('p:sp'), qn('p:grpSp'), qn('p:graphicFrame'), qn('p:cxnSp')):
            tree.append(copy.deepcopy(el))
        elif el.tag == qn('p:pic'):
            raise SystemExit('picture not supported')
    if s.has_notes_slide:
        src_tree = s.notes_slide.shapes._spTree
        dn = ns.notes_slide
        dst_tree = dn.shapes._spTree
        dst_tree.getparent().replace(dst_tree, copy.deepcopy(src_tree))
from script import SCRIPT
for n, txt in SCRIPT.items():
    sl = orig.slides[n-1]
    tf = sl.notes_slide.notes_text_frame
    assert tf is not None, n
    tf.text = txt
orig.save('merged.pptx')
print('slides', len(orig.slides))
